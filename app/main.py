"""AnavClouds Lead Puller — web app (FastAPI)."""
import csv
import hashlib
import hmac
import io
import json
import os
import secrets
import time
from contextlib import asynccontextmanager

from fastapi import Depends, FastAPI, File, HTTPException, Request, Response, UploadFile
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse

from . import claude, db, seamless, worker
from .config import cfg, shown_settings
from .filters import FILTERS, clean_filters, search_filter_count

STATIC = os.path.join(os.path.dirname(__file__), "static")
_worker = None


@asynccontextmanager
async def lifespan(app):
    global _worker
    db.connect()
    if not os.environ.get("DISABLE_WORKER"):
        _worker = worker.Worker()
        _worker.start()
    yield
    if _worker:
        _worker.stop()


app = FastAPI(title="AnavClouds Lead Puller", lifespan=lifespan, docs_url=None, redoc_url=None)


# ---------------------------------------------------------------- errors
@app.exception_handler(worker.PullError)
async def _pull_err(request, exc):
    return JSONResponse({"error": str(exc)}, status_code=400)


@app.exception_handler(seamless.SeamlessError)
async def _seam_err(request, exc):
    return JSONResponse({"error": str(exc), "code": exc.code}, status_code=400)


@app.exception_handler(claude.ClaudeError)
async def _claude_err(request, exc):
    return JSONResponse({"error": str(exc)}, status_code=400)


# ---------------------------------------------------------------- auth
def _check_password(password):
    return bool(cfg.APP_PASSWORD) and hmac.compare_digest(password.encode(), cfg.APP_PASSWORD.encode())


def require_login(request: Request):
    token = request.cookies.get("lp_session")
    if token:
        row = db.q1("SELECT created_at FROM sessions WHERE token=?", (token,))
        if row and time.time() - row["created_at"] < cfg.SESSION_DAYS * 86400:
            return True
    raise HTTPException(401, "Please log in.")


def _start_session(response):
    token = secrets.token_urlsafe(32)
    db.x("INSERT INTO sessions(token, created_at) VALUES(?,?)", (token, time.time()))
    response.set_cookie("lp_session", token, max_age=cfg.SESSION_DAYS * 86400, httponly=True, samesite="lax",
                        secure=cfg.COOKIE_SECURE)


_login_fails = []


@app.get("/api/auth/status")
def auth_status(request: Request):
    try:
        require_login(request)
        ok = True
    except HTTPException:
        ok = False
    return {"logged_in": ok, "no_password": not cfg.APP_PASSWORD}


@app.post("/api/auth/login")
async def auth_login(request: Request, response: Response):
    now = time.time()
    _login_fails[:] = [t for t in _login_fails if now - t < 300]
    if len(_login_fails) >= 10:
        raise HTTPException(429, "Too many wrong passwords. Wait 5 minutes.")
    body = await request.json()
    if not cfg.APP_PASSWORD:
        raise HTTPException(400, "No password set. Ask your developer to set APP_PASSWORD in .env.")
    if not _check_password(body.get("password") or ""):
        _login_fails.append(now)
        raise HTTPException(401, "Wrong password.")
    _start_session(response)
    return {"ok": True}


@app.post("/api/auth/logout")
def auth_logout(request: Request, response: Response):
    token = request.cookies.get("lp_session")
    if token:
        db.x("DELETE FROM sessions WHERE token=?", (token,))
    response.delete_cookie("lp_session")
    return {"ok": True}


# ---------------------------------------------------------------- meta & settings
def _mask(key):
    if not key:
        return ""
    return "•••• " + key[-4:]


def _credit_summary():
    budget = cfg.MONTHLY_CREDIT_BUDGET
    spent = db.month_spend()
    return {"balance": db.get_setting("credit_balance"), "balance_at": db.get_setting("credit_balance_at"),
            "monthly_budget": budget, "spent_this_month": spent,
            "budget_left": max(0, budget - spent) if budget else None}


@app.get("/api/meta", dependencies=[Depends(require_login)])
def meta():
    return {"filters": FILTERS, "credits": _credit_summary(),
            "has_seamless_key": bool(seamless.api_key()), "has_claude_key": bool(claude.api_key()),
            "defaults": {"per_company_cap": cfg.MAX_PER_COMPANY, "company_cooldown": cfg.COMPANY_COOLDOWN_DAYS,
                         "preview_size": cfg.PREVIEW_SIZE, "page_size": cfg.PAGE_SIZE}}


@app.get("/api/settings", dependencies=[Depends(require_login)])
def get_settings():
    """Read-only: every setting lives in .env. Change it there and restart the app."""
    return {"groups": shown_settings(), "credits": _credit_summary()}


@app.post("/api/settings/test-seamless", dependencies=[Depends(require_login)])
def test_seamless():
    return seamless.test_connection()


@app.post("/api/settings/test-claude", dependencies=[Depends(require_login)])
def test_claude():
    out = claude.prompt_to_filters("CTOs at software companies in the United States")
    return {"ok": True, "sample": out}


@app.get("/api/ledger", dependencies=[Depends(require_login)])
def ledger(limit: int = 200):
    rows = db.q("""SELECT l.*, p.name AS pull_name FROM ledger l LEFT JOIN pulls p ON p.id=l.pull_id
                   ORDER BY l.id DESC LIMIT ?""", (min(limit, 1000),))
    return {"rows": rows, "credits": _credit_summary()}


# ---------------------------------------------------------------- pulls
def _pull_stats(pid):
    one = lambda sql: db.q1(sql, (pid,))["n"]
    stats = {
        "seen": one("SELECT COUNT(*) n FROM contacts WHERE pull_id=?"),
        "kept": one("SELECT COUNT(*) n FROM contacts WHERE pull_id=? AND kept=1"),
        "skipped": one("SELECT COUNT(*) n FROM contacts WHERE pull_id=? AND kept=0"),
        "selected": one("SELECT COUNT(*) n FROM contacts WHERE pull_id=? AND kept=1 AND selected=1"),
        "to_research": one("""SELECT COUNT(*) n FROM contacts WHERE pull_id=? AND kept=1 AND selected=1
                              AND submitted_at IS NULL"""),
        "submitted": one("SELECT COUNT(*) n FROM contacts WHERE pull_id=? AND submitted_at IS NOT NULL"),
        "pending": one("SELECT COUNT(*) n FROM research_requests WHERE pull_id=? AND resolved_at IS NULL"),
        "with_email": one("""SELECT COUNT(*) n FROM contacts WHERE pull_id=? AND submitted_at IS NOT NULL
                             AND email IS NOT NULL AND email!=''"""),
        "fit": {r["fit"] or "unrated": r["n"] for r in db.q(
            "SELECT fit, COUNT(*) n FROM contacts WHERE pull_id=? AND kept=1 GROUP BY fit", (pid,))},
        "skip_reasons": db.q("""SELECT skip_reason AS reason, COUNT(*) n FROM contacts WHERE pull_id=? AND kept=0
                                GROUP BY skip_reason ORDER BY n DESC""", (pid,)),
        "by_country": db.q("""SELECT COALESCE(NULLIF(country,''),'Unknown') AS country, COUNT(*) n
                              FROM contacts WHERE pull_id=? AND kept=1 GROUP BY 1 ORDER BY n DESC LIMIT 12""", (pid,)),
        "by_size": db.q("""SELECT COALESCE(NULLIF(size,''),'Unknown') AS size, COUNT(*) n
                           FROM contacts WHERE pull_id=? AND kept=1 GROUP BY 1 ORDER BY n DESC""", (pid,)),
        "hit_rate": db.q("""SELECT COALESCE(NULLIF(country,''),'Unknown') AS country, COUNT(*) researched,
                            SUM(CASE WHEN email IS NOT NULL AND email!='' THEN 1 ELSE 0 END) AS with_email
                            FROM contacts WHERE pull_id=? AND submitted_at IS NOT NULL
                            AND research_status NOT IN ('queued','researching') GROUP BY 1 ORDER BY researched DESC""",
                         (pid,)),
        "dup_missing": one("""SELECT COUNT(*) n FROM contacts WHERE pull_id=? AND submitted_at IS NOT NULL
                              AND LOWER(research_status) IN ('duplicate','contact-already-researched')
                              AND (email IS NULL OR email='')"""),
        "research_outcomes": db.q("""SELECT research_status AS status, COUNT(*) n FROM contacts
                                     WHERE pull_id=? AND submitted_at IS NOT NULL GROUP BY 1 ORDER BY n DESC""", (pid,)),
    }
    return stats


def _pull_out(pid):
    p = db.get_pull(pid)
    if not p:
        raise HTTPException(404, "Pull not found.")
    p["stats"] = _pull_stats(pid)
    p["credits"] = _credit_summary()
    return p


@app.get("/api/pulls", dependencies=[Depends(require_login)])
def list_pulls():
    rows = db.q("""SELECT p.id, p.name, p.status, p.target, p.search_credits, p.research_credits, p.created_at,
                   p.finished_at, p.total_matches, p.cloned_from,
                   (SELECT COUNT(*) FROM contacts c WHERE c.pull_id=p.id AND c.kept=1) AS kept,
                   (SELECT COUNT(*) FROM contacts c WHERE c.pull_id=p.id AND c.email IS NOT NULL AND c.email!='') AS with_email
                   FROM pulls p ORDER BY p.id DESC""")
    return {"rows": rows}


@app.post("/api/pulls/interpret", dependencies=[Depends(require_login)])
async def interpret(request: Request):
    body = await request.json()
    prompt = (body.get("prompt") or "").strip()
    if not prompt:
        raise HTTPException(400, "Describe who you want to find.")
    pid = body.get("pull_id")
    if pid:
        pull = db.get_pull(pid)
        if pull["status"] not in ("draft", "previewed", "collected"):
            raise worker.PullError("Filters can't change while the pull is running.")
        out = claude.prompt_to_filters(prompt, pull["filters"])
        db.update_pull(pid, filters=out["filters"], notes=out["notes"], status="draft",
                       prompt=(pull["prompt"] or "") + "\n" + prompt)
        db.x("DELETE FROM contacts WHERE pull_id=? AND submitted_at IS NULL", (pid,))
        return _pull_out(pid)
    out = claude.prompt_to_filters(prompt)
    pid = db.create_pull(out["name"], prompt, out["filters"], out["notes"], target=out["target"])
    return _pull_out(pid)


@app.post("/api/pulls", dependencies=[Depends(require_login)])
async def create_pull(request: Request):
    body = await request.json()
    filters = clean_filters(body.get("filters") or {})
    pid = db.create_pull((body.get("name") or "New pull")[:80], body.get("prompt") or "", filters, [],
                         target=int(body.get("target") or 0))
    return _pull_out(pid)


@app.get("/api/pulls/{pid}", dependencies=[Depends(require_login)])
def get_pull(pid: int):
    return _pull_out(pid)


@app.patch("/api/pulls/{pid}", dependencies=[Depends(require_login)])
async def edit_pull(pid: int, request: Request):
    body = await request.json()
    pull = db.get_pull(pid)
    if not pull:
        raise HTTPException(404)
    fields = {}
    if "name" in body:
        fields["name"] = (body["name"] or "Untitled")[:80]
    for k in ("company_cooldown", "per_company_cap"):
        if k in body and body[k] not in (None, ""):
            if pull["status"] != "draft":
                raise worker.PullError("Change this before Preview (or use “Edit filters”).")
            fields[k] = max(0, int(body[k]))
    if "filters" in body:
        if pull["status"] not in ("draft", "previewed", "collected"):
            raise worker.PullError("Filters can't change while the pull is running.")
        if pull["status"] == "collected" and _pull_stats(pid)["submitted"]:
            raise worker.PullError("Research has started on this pull. Use “Run again” to change the filters.")
        fields["filters"] = clean_filters(body["filters"])
        fields["status"] = "draft"
        db.x("DELETE FROM contacts WHERE pull_id=?", (pid,))
    if fields:
        db.update_pull(pid, **fields)
    return _pull_out(pid)


@app.post("/api/pulls/{pid}/preview", dependencies=[Depends(require_login)])
async def preview(pid: int, request: Request):
    try:
        body = await request.json()
    except Exception:
        body = {}
    worker.run_preview(pid, from_top=bool(body.get("from_top")))
    return _pull_out(pid)


@app.post("/api/pulls/{pid}/collect", dependencies=[Depends(require_login)])
async def collect(pid: int, request: Request):
    body = await request.json()
    worker.start_collect(pid, body.get("target") or 0, body.get("per_company_cap") or 0)
    return _pull_out(pid)


@app.post("/api/pulls/{pid}/research", dependencies=[Depends(require_login)])
async def research(pid: int, request: Request):
    body = await request.json()
    worker.confirm_research(pid, body.get("choice") or "all", body.get("first_n"))
    return _pull_out(pid)


@app.post("/api/pulls/{pid}/fill-saved", dependencies=[Depends(require_login)])
def fill_saved(pid: int):
    pull = db.get_pull(pid)
    if not pull:
        raise HTTPException(404)
    if pull["status"] in ("collecting", "researching"):
        raise worker.PullError("Wait until research has finished.")
    result = worker.fill_from_saved(pid)
    out = _pull_out(pid)
    out["fill_result"] = result
    return out


@app.post("/api/pulls/{pid}/{action}", dependencies=[Depends(require_login)])
def pull_action(pid: int, action: str):
    if action == "pause":
        worker.pause(pid)
    elif action == "resume":
        worker.resume(pid)
    elif action == "cancel":
        worker.cancel(pid)
    elif action == "rerun":
        pull = db.get_pull(pid)
        new_id = db.create_pull(pull["name"] + " (new people)", pull["prompt"], pull["filters"], pull["notes"],
                                target=pull["target"], per_company_cap=pull["per_company_cap"], cloned_from=pid,
                                cooldown=pull["company_cooldown"])
        return _pull_out(new_id)
    elif action == "back-to-filters":
        pull = db.get_pull(pid)
        if pull["status"] in ("previewed",):
            db.update_pull(pid, status="draft")
    else:
        raise HTTPException(404)
    return _pull_out(pid)


@app.delete("/api/pulls/{pid}", dependencies=[Depends(require_login)])
def delete_pull(pid: int):
    pull = db.get_pull(pid)
    if not pull:
        raise HTTPException(404)
    if pull["status"] in ("collecting", "researching"):
        raise worker.PullError("Pause or stop the pull before deleting it.")
    # “seen” keys stay, so people you already paid for are never pulled again.
    db.x("DELETE FROM contacts WHERE pull_id=?", (pid,))
    db.x("DELETE FROM research_requests WHERE pull_id=?", (pid,))
    db.x("DELETE FROM pulls WHERE id=?", (pid,))
    return {"ok": True}


CONTACT_COLS = """id, name, first_name, last_name, title, company, domain, country, state, city, li_url, seniority,
                  department, size, industry, fit, fit_reason, kept, skip_reason, selected, research_status,
                  research_message, submitted_at, email, email_status, email_confidence, phone, phone_type"""


@app.get("/api/pulls/{pid}/contacts", dependencies=[Depends(require_login)])
def pull_contacts(pid: int, view: str = "kept", offset: int = 0, limit: int = 100, fit: str = "", q: str = ""):
    where = ["pull_id=?"]
    args = [pid]
    if view == "kept":
        where.append("kept=1")
    elif view == "skipped":
        where.append("kept=0")
    elif view == "results":
        where.append("submitted_at IS NOT NULL")
    elif view == "with_email":
        where.append("submitted_at IS NOT NULL AND email IS NOT NULL AND email!=''")
    elif view == "no_email":
        where.append("submitted_at IS NOT NULL AND (email IS NULL OR email='') "
                     "AND research_status NOT IN ('queued','researching')")
    if fit:
        where.append("COALESCE(fit,'unrated')=?")
        args.append(fit)
    if q:
        where.append("(name LIKE ? OR company LIKE ? OR title LIKE ? OR email LIKE ?)")
        args += [f"%{q}%"] * 4
    w = " AND ".join(where)
    total = db.q1(f"SELECT COUNT(*) n FROM contacts WHERE {w}", args)["n"]
    rows = db.q(f"SELECT {CONTACT_COLS} FROM contacts WHERE {w} ORDER BY id LIMIT ? OFFSET ?",
                (*args, min(limit, 500), max(offset, 0)))
    return {"total": total, "rows": rows}


@app.post("/api/pulls/{pid}/contacts/select", dependencies=[Depends(require_login)])
async def select_contacts(pid: int, request: Request):
    body = await request.json()
    ids = [int(i) for i in body.get("ids") or []]
    sel = 1 if body.get("selected") else 0
    pull = db.get_pull(pid)
    if pull["status"] not in ("collected", "previewed"):
        raise worker.PullError("Selection can only change before research starts.")
    for cid in ids:
        db.x("UPDATE contacts SET selected=? WHERE id=? AND pull_id=? AND kept=1 AND submitted_at IS NULL",
             (sel, cid, pid))
    return _pull_out(pid)


# ---------------------------------------------------------------- export
FRIENDLY_STATUS = {
    "done": "Found", "not found": "Not found in Seamless", "missing": "Not found in Seamless",
    "duplicate": "Already researched in your Seamless account before (no new charge)",
    "contact-already-researched": "Already researched in your Seamless account before (no new charge)",
    "error": "Seamless error", "no license or credits available": "No credits / license at the time",
    "timed out": "Seamless took too long", "no answer": "Seamless returned no result",
}


def _export_rows(pid, only_email):
    rows = db.q("""SELECT * FROM contacts WHERE pull_id=? AND submitted_at IS NOT NULL ORDER BY id""", (pid,))
    excl = {r["key"] for r in db.q("SELECT key FROM exclusions WHERE kind IN ('email','domain')")}
    out = []
    for r in rows:
        email = (r["email"] or "").strip()
        if only_email and not email:
            continue
        if email and ("em:" + email.lower() in excl or "domain:" + email.split("@")[-1].lower() in excl):
            continue
        out.append(r)
    return out


def _cell(v):
    """Stop spreadsheet formula injection: text starting with = + - @ (or tab/CR) is prefixed with '."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


def _risky(status):
    s = (status or "").lower()
    return s and s not in ("valid", "verified", "deliverable")


@app.get("/api/pulls/{pid}/export", dependencies=[Depends(require_login)])
def export(pid: int, format: str = "anavmail", include_risky: int = 1):
    pull = db.get_pull(pid)
    if not pull:
        raise HTTPException(404)
    safe = "".join(ch if ch.isalnum() else "-" for ch in pull["name"]).strip("-")[:40] or f"pull-{pid}"
    if format == "anavmail":
        rows = [r for r in _export_rows(pid, True)
                if include_risky or not _risky(r["email_status"])]
        rows = [r for r in rows if (r["email_status"] or "").lower() != "invalid"]
        seen_emails, unique = set(), []
        for r in rows:
            e = r["email"].strip().lower()
            if e not in seen_emails:
                seen_emails.add(e)
                unique.append(r)
        rows = unique
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(["email", "first_name", "last_name", "company", "title", "linkedin"])
        for r in rows:
            w.writerow([_cell(r[k]) for k in ("email", "first_name", "last_name", "company", "title", "li_url")])
        return _download(buf.getvalue().encode("utf-8-sig"), f"{safe}-anavmail.csv", "text/csv")
    cols = [("first_name", "First name"), ("last_name", "Last name"), ("email", "Email"),
            ("email_status", "Email status"), ("email_confidence", "Email confidence"),
            ("personal_email", "Personal email"), ("phone", "Phone"), ("phone_type", "Phone type"),
            ("company_phone", "Company phone"), ("title", "Title"), ("seniority", "Seniority"),
            ("department", "Department"), ("company", "Company"), ("domain", "Domain"), ("industry", "Industry"),
            ("size", "Company size"), ("city", "City"), ("state", "State"), ("country", "Country"),
            ("li_url", "LinkedIn"), ("fit", "Fit"), ("fit_reason", "Fit note"), ("research_status", "Result")]
    rows = _export_rows(pid, False)

    def val(r, k):
        if k == "research_status":
            return FRIENDLY_STATUS.get((r[k] or "").lower(), r[k])
        return _cell(r[k])

    if format == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font
        wb = Workbook()
        ws = wb.active
        ws.title = "Contacts"
        ws.append([c[1] for c in cols])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for r in rows:
            ws.append([val(r, c[0]) for c in cols])
        ws.freeze_panes = "A2"
        for i, (k, _) in enumerate(cols, start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = 28 if k in ("email", "li_url", "company", "title") else 16
        b = io.BytesIO()
        wb.save(b)
        return _download(b.getvalue(), f"{safe}.xlsx",
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([c[1] for c in cols])
    for r in rows:
        w.writerow([val(r, c[0]) for c in cols])
    return _download(buf.getvalue().encode("utf-8-sig"), f"{safe}-full.csv", "text/csv")


def _download(data, filename, mime):
    return StreamingResponse(iter([data]), media_type=mime,
                             headers={"Content-Disposition": f'attachment; filename="{filename}"'})


# ---------------------------------------------------------------- presets
@app.get("/api/presets", dependencies=[Depends(require_login)])
def presets():
    rows = db.q("SELECT * FROM presets ORDER BY name")
    for r in rows:
        r["filters"] = json.loads(r.pop("filters_json") or "{}")
    return {"rows": rows}


@app.post("/api/presets", dependencies=[Depends(require_login)])
async def save_preset(request: Request):
    body = await request.json()
    name = (body.get("name") or "").strip()
    if not name:
        raise HTTPException(400, "Give the preset a name.")
    filters = clean_filters(body.get("filters") or {})
    db.x("INSERT INTO presets(name,prompt,filters_json,created_at) VALUES(?,?,?,?)",
         (name[:80], body.get("prompt") or "", json.dumps(filters), time.time()))
    return presets()


@app.post("/api/presets/{preset_id}/use", dependencies=[Depends(require_login)])
def use_preset(preset_id: int):
    p = db.q1("SELECT * FROM presets WHERE id=?", (preset_id,))
    if not p:
        raise HTTPException(404)
    pid = db.create_pull(p["name"], p["prompt"], json.loads(p["filters_json"]), [])
    return _pull_out(pid)


@app.delete("/api/presets/{preset_id}", dependencies=[Depends(require_login)])
def delete_preset(preset_id: int):
    db.x("DELETE FROM presets WHERE id=?", (preset_id,))
    return presets()


# ---------------------------------------------------------------- exclusions (do-not-pull list)
def _add_exclusion(kind, value, source):
    value = (value or "").strip()
    if not value:
        return False
    if kind == "email":
        if "@" not in value:
            return False
        key = "em:" + value.lower()
    elif kind == "domain":
        value = db.norm_domain(value.split("@")[-1])
        if "." not in value:
            return False
        key = "domain:" + value
    elif kind == "linkedin":
        key = "li:" + db.norm_li(value)
    elif kind == "name_company":
        key = value
    else:
        return False
    db.x("INSERT OR IGNORE INTO exclusions(key,kind,value,source,created_at) VALUES(?,?,?,?,?)",
         (key, kind, value, source, time.time()))
    return True


@app.get("/api/exclusions", dependencies=[Depends(require_login)])
def list_exclusions(q: str = "", offset: int = 0, limit: int = 50):
    where, args = "", []
    if q:
        where, args = "WHERE value LIKE ?", [f"%{q}%"]
    total = db.q1(f"SELECT COUNT(*) n FROM exclusions {where}", args)["n"]
    rows = db.q(f"SELECT * FROM exclusions {where} ORDER BY created_at DESC LIMIT ? OFFSET ?",
                (*args, min(limit, 200), max(offset, 0)))
    counts = {r["kind"]: r["n"] for r in db.q("SELECT kind, COUNT(*) n FROM exclusions GROUP BY kind")}
    return {"total": total, "rows": rows, "counts": counts}


def _pick(header, *names):
    low = [h.strip().lower().replace(" ", "_") for h in header]
    for n in names:
        if n in low:
            return low.index(n)
    return None


@app.post("/api/exclusions/upload", dependencies=[Depends(require_login)])
async def upload_exclusions(file: UploadFile = File(...)):
    raw = (await file.read()).decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(raw))
    header = next(reader, [])
    ie = _pick(header, "email", "email_address", "e-mail")
    idom = _pick(header, "domain", "website", "company_domain")
    ili = _pick(header, "linkedin", "linkedin_url", "li_url", "linkedin_profile")
    ifn, iln = _pick(header, "first_name", "firstname", "first"), _pick(header, "last_name", "lastname", "last")
    iname = _pick(header, "name", "full_name")
    ico = _pick(header, "company", "company_name", "organization")
    if all(i is None for i in (ie, idom, ili, ico)):
        raise HTTPException(400, "Couldn't find an email, domain, linkedin or company column in that file.")
    added = 0
    src = file.filename or "upload"
    for row in reader:
        get = lambda i: row[i].strip() if i is not None and i < len(row) else ""
        if get(ie):
            added += _add_exclusion("email", get(ie), src)
        if get(idom) and not get(ie):
            added += _add_exclusion("domain", get(idom), src)
        if get(ili):
            added += _add_exclusion("linkedin", get(ili), src)
        name = get(iname) or f"{get(ifn)} {get(iln)}".strip()
        if name and get(ico):
            added += _add_exclusion("name_company", "nc:" + db.norm(name) + "|" + db.norm(get(ico)), src)
    return {"added": added, **list_exclusions()}


@app.post("/api/exclusions", dependencies=[Depends(require_login)])
async def add_exclusions(request: Request):
    body = await request.json()
    added = 0
    for line in (body.get("text") or "").replace(",", "\n").splitlines():
        line = line.strip()
        if not line:
            continue
        if "linkedin.com" in line.lower():
            added += _add_exclusion("linkedin", line, "typed")
        elif "@" in line and not line.startswith("@"):
            added += _add_exclusion("email", line, "typed")
        else:
            added += _add_exclusion("domain", line.lstrip("@"), "typed")
    return {"added": added, **list_exclusions()}


@app.delete("/api/exclusions", dependencies=[Depends(require_login)])
def clear_exclusions(key: str = ""):
    if key:
        db.x("DELETE FROM exclusions WHERE key=?", (key,))
    else:
        db.x("DELETE FROM exclusions")
    return list_exclusions()


# ---------------------------------------------------------------- company memory
@app.get("/api/companies", dependencies=[Depends(require_login)])
def list_companies(q: str = "", status: str = "all", offset: int = 0, limit: int = 50):
    days = cfg.COMPANY_COOLDOWN_DAYS
    now = time.time()
    where, args = ["1=1"], []
    if q:
        where.append("(name LIKE ? OR domain LIKE ?)")
        args += [f"%{q}%", f"%{q}%"]
    cool = "(last_at > ? AND (released_at IS NULL OR released_at < last_at))"
    if status == "cooldown":
        where.append(cool)
        args.append(now - days * 86400)
    elif status == "available":
        where.append("NOT " + cool)
        args.append(now - days * 86400)
    w = " AND ".join(where)
    total = db.q1(f"SELECT COUNT(*) n FROM companies WHERE {w}", args)["n"]
    rows = db.q(f"SELECT * FROM companies WHERE {w} ORDER BY last_at DESC LIMIT ? OFFSET ?",
                (*args, min(limit, 200), max(offset, 0)))
    blocked = {r["key"] for r in db.q("SELECT key FROM exclusions WHERE kind='domain'")}
    for r in rows:
        released = r["released_at"] and r["released_at"] >= r["last_at"]
        r["available_at"] = None if released else r["last_at"] + days * 86400
        r["on_cooldown"] = bool(days and not released and r["available_at"] > now)
        r["blocked"] = bool(r["domain"]) and ("domain:" + r["domain"]) in blocked
    counts = {
        "all": db.q1("SELECT COUNT(*) n FROM companies")["n"],
        "cooldown": db.q1(f"SELECT COUNT(*) n FROM companies WHERE {cool}", (now - days * 86400,))["n"] if days else 0,
    }
    return {"total": total, "rows": rows, "counts": counts, "cooldown_days": days}


@app.post("/api/companies/release", dependencies=[Depends(require_login)])
async def release_company(request: Request):
    body = await request.json()
    db.x("UPDATE companies SET released_at=? WHERE key=?", (time.time(), body.get("key")))
    return {"ok": True}


@app.post("/api/companies/block", dependencies=[Depends(require_login)])
async def block_company(request: Request):
    body = await request.json()
    row = db.q1("SELECT * FROM companies WHERE key=?", (body.get("key"),))
    if not row:
        raise HTTPException(404)
    if row["domain"]:
        _add_exclusion("domain", row["domain"], "company memory")
    else:
        raise worker.PullError("This company has no domain on record, so it can't be blocked by domain.")
    return {"ok": True}


# ---------------------------------------------------------------- page
@app.get("/")
def index():
    return FileResponse(os.path.join(STATIC, "index.html"), headers={"Cache-Control": "no-store"})


@app.get("/healthz")
def healthz():
    return {"ok": True}
