import io
import os
import re

from app import config, db, engine
from tests.conftest import US, check_books, run, start


def test_login_is_required(env):
    client, fake, clock = env
    client.post("/api/auth/logout")
    assert client.get("/api/meta").status_code == 401
    assert client.post("/api/auth/login", json={"password": "wrong"}).status_code == 401
    assert client.post("/api/auth/login", json={"password": "test-pass-123"}).status_code == 200
    assert client.get("/api/meta").status_code == 200


def test_settings_page_masks_secrets(env):
    client, fake, clock = env
    s = client.get("/api/settings").json()
    flat = {r["key"]: r["value"] for g in s["groups"] for r in g["rows"]}
    assert flat["APP_PASSWORD"] == "•••• -123" and flat["ANTHROPIC_API_KEY"].startswith("••••")
    meta = client.get("/api/meta").json()
    assert all("key-alpha-1111" not in str(a) for a in meta["overview"]["accounts"])


def test_every_setting_is_in_env_example():
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    example = open(os.path.join(root, ".env.example")).read()
    missing = [k for k, *_ in config.SETTINGS if not re.search(r"^#? ?" + k + "=", example, re.M)]
    assert not missing, missing


def test_env_file_is_read(tmp_path, monkeypatch):
    f = tmp_path / ".env"
    f.write_text('MAX_PER_COMPANY=7\nMOVE_ON_SEARCH_COST="0.5"\n# comment\nFRESH_RESCAN=0  # off\n')
    for k in ("MAX_PER_COMPANY", "MOVE_ON_SEARCH_COST", "FRESH_RESCAN"):
        monkeypatch.delenv(k, raising=False)
    try:
        config._load_dotenv(str(f))
        assert config.cfg.MAX_PER_COMPANY == 7 and config.cfg.MOVE_ON_SEARCH_COST == 0.5
        assert config.cfg.FRESH_RESCAN is False
    finally:
        for k in ("MAX_PER_COMPANY", "MOVE_ON_SEARCH_COST", "FRESH_RESCAN"):
            os.environ.pop(k, None)


def test_describe_fills_only_valid_filters(env):
    client, fake, clock = env
    r = client.post("/api/describe", json={"text": "CTOs in Israel, 1-500 employees, no interns"}).json()
    f = r["filters"]
    assert f["jobTitle"] == {"in": ["CTO", "Chief Technology Officer"], "out": ["intern"]}
    assert "Bogus" not in f["companySize"]["in"] and f["industry"]["in"] == ["Computer Software"]


def test_a_pull_needs_real_filters_and_a_target(env):
    client, fake, clock = env
    r = client.post("/api/pulls", json={"filters": {"technologiesIsOr": True}, "target": 10})
    assert r.status_code == 400 and "filter" in r.json()["error"]
    r = client.post("/api/pulls", json={"filters": US, "target": 0})
    assert r.status_code == 400


def test_only_one_pull_runs_at_a_time(env):
    client, fake, clock = env
    a = start(client, US, target=40, max_credits=400)
    b = start(client, {"contactCountry": {"in": ["Canada"]}}, target=20, max_credits=400)
    engine.tick()
    engine.tick()
    assert client.get(f"/api/pulls/{b}").json()["status"] == "queued"
    run(client, a)
    run(client, b)
    check_books(fake)


def test_downloads(env):
    client, fake, clock = env
    us = [p for p in fake.people if p["country"] == "United States" and p["domain"]][:30]
    fake.email_override[us[0]["pid"]] = "=HYPERLINK(\"x\")@evil.com"
    pid = start(client, US, target=60, max_credits=500)
    run(client, pid)
    text = client.get(f"/api/pulls/{pid}/download").content.decode("utf-8-sig")
    lines = text.strip().splitlines()
    assert lines[0] == "email,first_name,last_name,company"
    for line in lines[1:]:
        assert not line.startswith("="), "a cell would run as a spreadsheet formula"
    invalid = [c["email"] for c in db.q("SELECT email FROM contacts WHERE pull_id=? AND email_status='invalid'", (pid,))]
    assert not any(e in text for e in invalid)
    custom = client.get(f"/api/pulls/{pid}/download?columns=company,title").content.decode("utf-8-sig")
    assert custom.splitlines()[0] == "email,company,title"
    x = client.get(f"/api/pulls/{pid}/download?kind=all&format=xlsx")
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.active.max_row > 10
    ne = client.get(f"/api/pulls/{pid}/download?kind=no_email").content.decode("utf-8-sig")
    assert ne.splitlines()[0].startswith("full_name")


def test_people_pages_and_search(env):
    client, fake, clock = env
    pid = start(client, US, target=80, max_credits=600)
    run(client, pid)
    p1 = client.get(f"/api/pulls/{pid}/people?view=emails&limit=25&offset=0").json()
    p2 = client.get(f"/api/pulls/{pid}/people?view=emails&limit=25&offset=25").json()
    assert p1["total"] >= 80 and len(p1["rows"]) == 25
    assert not {r["id"] for r in p1["rows"]} & {r["id"] for r in p2["rows"]}
    name = p1["rows"][0]["name"].split()[0]
    hits = client.get(f"/api/pulls/{pid}/people?view=emails&q={name}").json()
    assert hits["total"] >= 1 and all(name in r["name"] for r in hits["rows"])
    sk = client.get(f"/api/pulls/{pid}/people?view=skipped").json()
    assert all(r["skip_reason"] for r in sk["rows"])
    lst = client.get("/api/pulls?limit=1").json()
    assert lst["total"] == 1 and lst["rows"][0]["emails"] >= 80


def test_place_and_technology_suggestions(env):
    client, fake, clock = env
    assert client.get("/api/lookup/places?q=isr").json()["rows"][0]["value"] == "Israel"
    pid = start(client, US, target=40, max_credits=400)
    run(client, pid)
    rows = client.get("/api/lookup/tech?q=sales").json()["rows"]
    assert rows[0]["value"] == "Salesforce" and rows[0]["seen"] > 0     # learned from research results


def test_rerun_and_delete_keep_memory(env):
    client, fake, clock = env
    pid = start(client, US, target=30, max_credits=300)
    run(client, pid)
    paid = db.q1("SELECT COUNT(*) n FROM contacts WHERE pull_id=? AND submitted_at IS NOT NULL", (pid,))["n"]
    assert client.delete(f"/api/pulls/{pid}").status_code == 200
    again = client.post(f"/api/pulls/{pid}/rerun")
    assert again.status_code == 404
    pid2 = start(client, US, target=30, max_credits=300)
    run(client, pid2)
    from tests.conftest import check_people
    check_people(fake)
    assert paid > 0
    r = client.post(f"/api/pulls/{pid2}/rerun").json()
    assert r["status"] == "queued" and r["filters"] == client.get(f"/api/pulls/{pid2}").json()["filters"]


def test_do_not_pull_upload(env):
    client, fake, clock = env
    csv_text = "Email,Company,Website\na@x.com,,\n,Infosys,\n,,acme1.com\n"
    r = client.post("/api/exclusions/upload", files={"file": ("list.csv", csv_text, "text/csv")}).json()
    assert r["added"] == 3 and r["counts"] == {"email": 1, "company": 1, "domain": 1}
    r = client.get("/api/exclusions?q=infosys").json()
    assert r["total"] == 1


def test_download_never_repeats_an_email_and_never_runs_formulas(env):
    from app import exports
    client, fake, clock = env
    pid = engine.create_pull("x", "", {"contactCountry": {"in": ["Canada"]}}, 5, 50)
    for i, (email, company) in enumerate((("a@x.com", "=cmd|' /C calc'!A0"), ("A@x.com", "+SUM(1)"), ("b@x.com", "@evil"))):
        db.x("INSERT INTO contacts(pull_id,state,name,company,email,email_status,dup_email,submitted_at) "
             "VALUES(?,?,?,?,?,?,0,1)", (pid, "done", f"Person {i}", company, email, "valid"))
    rows = exports.anavmail_rows(pid)
    assert len(rows) == 2                                   # a@x.com and A@x.com are the same inbox
    text = exports.to_csv(rows, ["email", "company"]).decode("utf-8-sig")
    for line in text.splitlines()[1:]:
        assert line.split(",", 1)[1][:1] not in ("=", "+", "-", "@"), line
    assert exports.cell("=1+1") == "'=1+1" and exports.cell("Acme") == "Acme"


def test_broken_server_dns_never_skips_anyone(env, monkeypatch):
    from app import dnscheck
    monkeypatch.setattr(dnscheck, "LOOKUP", None)
    monkeypatch.setattr(dnscheck, "_health", {"ok": None, "at": 0.0})
    monkeypatch.setattr(dnscheck, "_real_lookup", lambda d: ("no", "domain does not exist"))
    assert dnscheck.has_mail_server("realcompany.com") == "unknown"     # DNS says nothing exists: don't trust it
    monkeypatch.setattr(dnscheck, "_health", {"ok": None, "at": 0.0})
    monkeypatch.setattr(dnscheck, "_real_lookup",
                        lambda d: ("yes", "mx") if d in dnscheck.KNOWN_GOOD else ("no", "no MX"))
    assert dnscheck.has_mail_server("parked-domain.com") == "no"
